#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
audit_docs.py — 백서 · 관제탑(대시보드) · 세션정리 간 숫자 교차대조

목적: "대시보드가 옛날 숫자를 들고 있다"를 사람 눈 대신 코드가 잡는다.
LLM 상호검토와 달리 **문서에서 숫자를 직접 추출해 문자 단위로 대조**한다.

사용:
  python audit_docs.py --docs 백서_빌더/out/진우퀀트_백서.html \
                              진행현황/진우퀀트_진행현황.html \
                              감사/진우퀀트_세션정리_2026-07-27.md \
                       [--baseline 기준값.json]

동작:
  1. 각 문서에서 (지표명, 값) 쌍을 추출 — 앵커 사전 기반
  2. 문서 간 동일 지표 값이 다르면 🚨 (= 갱신 누락)
  3. baseline.json이 있으면 소스오브트루스와도 대조
"""
try:  # 윈도 파이프(cp949)에서 이모지 출력 시 UnicodeEncodeError 방지 (2026-07-27)
    import sys as _sys; _sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

import argparse, json, re, sys
from pathlib import Path
from collections import defaultdict

# 지표 앵커 — 세션정리 기준 핵심 숫자. 필요시 추가만 하면 됨.
ANCHORS = {
    "PBO":            r"PBO[^\d%]{0,12}([\d.]+)\s*%",
    "DSR":            r"DSR[^\d%]{0,12}([\d.]+)\s*%",
    "배당프리미엄":    r"배당[^\n]{0,20}?([+\-]?[\d.]+)\s*%\s*/?\s*yr",
    "BP프리미엄":      r"B/P[^\n]{0,20}?([+\-]?[\d.]+)\s*%",
    "주도주규칙_CAGR": r"8\.1\s*%|주도주[^\n]{0,30}?([\d.]+)\s*%",
    "KOSPI_2015_CAGR": r"KOSPI\s*\(?\s*([\d.]+)\s*%",
    "왕복비용":        r"왕복\s*([\d.]+)\s*%",
    "배당_HAC_t":     r"배당[^\n]{0,15}?t\s*[= ]\s*([\d.]+)",
    "BP_t":           r"B/P\s*([\d.]+)",
    "EP_t":           r"E/P\s*([\d.]+)",
    "CSCV_조합수":     r"([\d,]+)\s*조합",
    "격자설정수":      r"격자\s*([\d]+)\s*설정",
    "보유종목수":      r"top\s*(\d+)\s*동일가중",
    "유니버스":        r"top\s*(\d{3})",
    "KOSPI_검증_불일치": r"380\s*→\s*(\d+)",
    "잔여CA":         r"잔여\s*CA\s*([\d.]+)\s*%",
    "커버리지_2014이전": r"2014이전\s*커버리지\s*([\d.]+)\s*%",
}


def to_text(p: Path) -> str:
    suf = p.suffix.lower()
    raw = None
    if suf in (".html", ".htm"):
        raw = p.read_text(encoding="utf-8", errors="replace")
        raw = re.sub(r"<script.*?</script>|<style.*?</style>", " ", raw, flags=re.S | re.I)
        raw = re.sub(r"<[^>]+>", " ", raw)
    elif suf == ".md" or suf == ".txt":
        raw = p.read_text(encoding="utf-8", errors="replace")
    elif suf == ".docx":
        try:
            from docx import Document
            raw = "\n".join(x.text for x in Document(str(p)).paragraphs)
            for t in Document(str(p)).tables:
                for row in t.rows:
                    raw += "\n" + " | ".join(c.text for c in row.cells)
        except ImportError:
            return "__NEED_python-docx__"
    elif suf == ".pdf":
        try:
            import pdfplumber
            with pdfplumber.open(str(p)) as pdf:
                raw = "\n".join((pg.extract_text() or "") for pg in pdf.pages)
        except ImportError:
            return "__NEED_pdfplumber__"
    else:
        return ""
    return re.sub(r"[ \t\u00a0]+", " ", raw or "")


def extract(text):
    out = {}
    for key, pat in ANCHORS.items():
        vals = [g for m in re.finditer(pat, text) for g in m.groups() if g]
        if vals:
            uniq = sorted(set(v.replace(",", "") for v in vals))
            out[key] = uniq
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--docs", nargs="+", required=True)
    ap.add_argument("--baseline", default=None)
    ap.add_argument("--out", default="감사_문서정합성.md")
    a = ap.parse_args()

    L = []
    def say(s=""):
        print(s); L.append(s)

    say("# 문서 정합성 감사 (백서 · 관제탑 · 세션정리)\n")

    found = {}
    for d in a.docs:
        p = Path(d)
        if not p.exists():
            say(f"❌ 파일 없음: {d}")
            continue
        t = to_text(p)
        if t.startswith("__NEED_"):
            say(f"⚠️ {p.name}: `pip install {t.replace('__NEED_','').replace('__','')}` 필요")
            continue
        found[p.name] = extract(t)
        say(f"- 읽음: `{p.name}` ({len(t):,}자, 지표 {len(found[p.name])}개 추출)")

    if len(found) < 2:
        say("\n비교할 문서가 2개 미만이다.")
        Path(a.out).write_text("\n".join(L), encoding="utf-8"); return

    # 지표별 문서간 대조
    say("\n## 지표별 교차대조\n")
    say("| 지표 | " + " | ".join(found) + " | 판정 |")
    say("|---|" + "---|" * (len(found) + 1))
    conflicts = []
    for key in ANCHORS:
        row = []
        vals_per_doc = []
        for doc in found:
            v = found[doc].get(key)
            row.append(", ".join(v) if v else "—")
            if v:
                vals_per_doc.append(set(v))
        if not vals_per_doc:
            continue
        if len(vals_per_doc) >= 2:
            inter = set.intersection(*vals_per_doc)
            verdict = "✅ 일치" if inter else "🚨 **불일치**"
            if not inter:
                conflicts.append((key, row))
        else:
            verdict = "· 단일문서"
        say(f"| {key} | " + " | ".join(row) + f" | {verdict} |")

    say("\n## 판정\n")
    if conflicts:
        say(f"🚨 **불일치 {len(conflicts)}건 — 갱신 누락 의심**\n")
        for k, r in conflicts:
            say(f"- `{k}`: {r}")
        say("\n→ 어느 쪽이 최신인지 확인하고, **오래된 문서를 재생성**할 것.")
        say("  대시보드는 백서에서 파생되어야 하며 손으로 숫자를 넣으면 이 불일치가 반복된다.")
    else:
        say("✅ 추출된 지표 전부 문서 간 일치.")

    # baseline 대조
    if a.baseline and Path(a.baseline).exists():
        base = json.loads(Path(a.baseline).read_text(encoding="utf-8"))
        say("\n## 소스오브트루스(baseline) 대조\n")
        for k, truth in base.items():
            for doc, ex in found.items():
                if k in ex and str(truth) not in ex[k]:
                    say(f"- 🚨 `{doc}` {k}: 문서 {ex[k]} vs 기준 {truth}")
        say("(출력 없으면 전부 일치)")
    else:
        say("\n> baseline 미지정. 파이프라인이 `기준값.json`을 뱉게 만들면")
        say("> 문서↔문서가 아니라 **문서↔산출물** 대조까지 가능하다. 그게 최종형이다.")

    Path(a.out).write_text("\n".join(L), encoding="utf-8")
    print(f"\n저장: {a.out}")
    sys.exit(1 if conflicts else 0)


if __name__ == "__main__":
    main()
